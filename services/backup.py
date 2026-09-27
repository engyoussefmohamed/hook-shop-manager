import os
import sqlite3
import threading
import time
from datetime import datetime

from database.connection import DB_PATH
from database.settings import get_settings, save_setting

_scheduler_started = False
_backup_lock = threading.Lock()


COLUMN_NAMES = {
    "products": {
        "id": "رقم", "name": "الاسم", "category": "الفئة",
        "price": "سعر البيع", "cost_price": "سعر التكلفة",
        "quantity": "الكمية", "barcode": "الباركود",
        "created_at": "تاريخ الإضافة", "min_stock": "الحد الأدنى",
    },
    "categories": {
        "id": "رقم", "name": "الفئة",
    },
    "sales": {
        "id": "رقم", "total": "الإجمالي", "discount": "الخصم",
        "final_total": "الإجمالي النهائي", "payment_method": "طريقة الدفع",
        "notes": "ملاحظات", "created_at": "التاريخ", "voided": "ملغاة",
        "paid_amount": "المدفوع", "remaining_amount": "المتبقي",
        "tax": "ضريبة", "service": "خدمة",
        "discount_type": "نوع الخصم", "discount_value": "قيمة الخصم",
    },
    "sale_items": {
        "id": "رقم", "sale_id": "رقم الفاتورة", "product_id": "رقم المنتج",
        "product_name": "اسم المنتج", "quantity": "الكمية",
        "unit_price": "سعر الوحدة", "cost_price": "سعر التكلفة",
        "total_price": "الإجمالي",
    },
    "returns": {
        "id": "رقم", "sale_id": "رقم الفاتورة", "product_id": "رقم المنتج",
        "product_name": "اسم المنتج", "quantity": "الكمية", "amount": "المبلغ",
        "return_type": "نوع المرتجع", "exchange_product_id": "رقم المنتج المستبدل",
        "notes": "ملاحظات", "affected_balance": "رصيد متأثر", "created_at": "التاريخ",
    },
    "stock_movements": {
        "id": "رقم", "product_id": "رقم المنتج", "product_name": "اسم المنتج",
        "movement_type": "نوع الحركة", "quantity": "الكمية",
        "old_quantity": "قبل", "new_quantity": "بعد", "notes": "ملاحظات",
        "ref_type": "نوع المرجع", "ref_id": "رقم المرجع",
        "created_at": "التاريخ",
    },
    "settings": {
        "key": "المفتاح", "value": "القيمة",
    },
}


def _export_excel(folder, timestamp):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from database.connection import get_db

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    HEADER_FILL = PatternFill("solid", fgColor="1A1F36")

    tables = [
        ("المنتجات", "products"),
        ("الفئات", "categories"),
        ("المبيعات", "sales"),
        ("بنود المبيعات", "sale_items"),
        ("المرتجعات", "returns"),
        ("حركات المخزون", "stock_movements"),
        ("الإعدادات", "settings"),
    ]

    with get_db() as conn:
        for sheet_name, table in tables:
            rows = conn.execute(f"SELECT * FROM {table}").fetchall()
            ws = wb.create_sheet(title=sheet_name)
            if not rows:
                ws.append([sheet_name])
                continue
            col_map = COLUMN_NAMES.get(table, {})
            db_cols = [c for c in rows[0].keys() if c in col_map]
            headers = [col_map[c] for c in db_cols]
            ws.append(headers)
            for cell in ws[1]:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
            for row in rows:
                ws.append([row[h] for h in db_cols])

    fname = os.path.join(folder, f"shop_{timestamp}.xlsx")
    wb.save(fname)


_SKIP_SETTINGS = {"printer_name", "backup_path", "last_backup"}

_SHEET_TO_TABLE = {
    "المنتجات":      "products",
    "الفئات":        "categories",
    "المبيعات":      "sales",
    "بنود المبيعات": "sale_items",
    "المرتجعات":     "returns",
    "حركات المخزون": "stock_movements",
    "الإعدادات":     "settings",
}

_RESTORE_ORDER = [
    "الإعدادات",
    "الفئات",
    "المنتجات",
    "المبيعات",
    "بنود المبيعات",
    "المرتجعات",
    "حركات المخزون",
]

_TABLE_LABELS = {
    "products":        "منتج",
    "categories":      "فئة",
    "sales":           "فاتورة",
    "sale_items":      "بند بيع",
    "returns":         "مرتجع",
    "stock_movements": "حركة مخزون",
    "settings":        "إعداد",
}


def restore_from_excel(file_path):
    """
    Restore data from a backup Excel (.xlsx) file produced by _export_excel().
    Returns (counts_dict, errors_list).
    Machine-specific settings (printer_name, backup_path, last_backup) are skipped.
    """
    import openpyxl
    from database.connection import get_db

    # Build Arabic-header → DB-column reverse maps per table
    reverse_maps = {
        table: {arabic: eng for eng, arabic in col_map.items()}
        for table, col_map in COLUMN_NAMES.items()
    }

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    counts = {}
    errors = []

    with get_db() as conn:
        conn.execute("PRAGMA foreign_keys = OFF")

        for sheet_name in _RESTORE_ORDER:
            if sheet_name not in wb.sheetnames:
                continue

            table   = _SHEET_TO_TABLE[sheet_name]
            rev_map = reverse_maps.get(table, {})
            ws      = wb[sheet_name]

            all_rows = list(ws.iter_rows(values_only=True))
            if len(all_rows) < 2:
                counts[table] = 0
                continue

            # Map header positions to DB column names
            col_mapping = []
            for i, cell in enumerate(all_rows[0]):
                db_col = rev_map.get(str(cell) if cell is not None else "")
                if db_col:
                    col_mapping.append((i, db_col))

            if not col_mapping:
                counts[table] = 0
                continue

            count = 0
            for row_values in all_rows[1:]:
                if all(v is None for v in row_values):
                    continue

                row_data = {
                    col: (row_values[idx] if idx < len(row_values) else None)
                    for idx, col in col_mapping
                }

                # Skip machine-specific settings
                if table == "settings" and row_data.get("key") in _SKIP_SETTINGS:
                    continue

                cols_sql    = ", ".join(row_data.keys())
                placeholders = ", ".join(["?"] * len(row_data))
                try:
                    conn.execute(
                        f"INSERT OR REPLACE INTO {table} ({cols_sql}) VALUES ({placeholders})",
                        list(row_data.values()),
                    )
                    count += 1
                except Exception as exc:
                    errors.append(f"{table}: {exc}")

            counts[table] = count

        conn.execute("PRAGMA foreign_keys = ON")

    wb.close()
    return counts, errors


def find_google_drive_path():
    candidates = [
        os.path.expanduser("~/Google Drive"),
        os.path.expanduser("~/My Drive"),
        os.path.expanduser("~/OneDrive"),
        os.path.expanduser("~/Dropbox"),
    ]
    try:
        import winreg
        try:
            key = winreg.OpenKey(
                winreg.HKEY_CURRENT_USER, r"Software\Google\DriveFS"
            )
            try:
                path, _ = winreg.QueryValueEx(key, "PerAccountPreferences")
                if path and isinstance(path, str) and os.path.exists(path):
                    return path
            finally:
                winreg.CloseKey(key)
        except Exception:
            pass
    except Exception:
        pass

    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def _sqlite_backup(src_path, dest_path):
    src = sqlite3.connect(src_path)
    dst = sqlite3.connect(dest_path)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()


def do_backup():
    settings    = get_settings()
    backup_path = settings.get("backup_path", "").strip()

    if not backup_path:
        backup_path = find_google_drive_path()

    if not backup_path or not os.path.exists(backup_path):
        backup_path = os.path.dirname(DB_PATH)

    try:
        folder = os.path.join(backup_path, "shop_backup")
        try:
            os.makedirs(folder, exist_ok=True)
        except PermissionError:
            folder = os.path.join(os.path.dirname(DB_PATH), "shop_backup")
            os.makedirs(folder, exist_ok=True)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest      = os.path.join(folder, f"shop_{timestamp}.db")
        try:
            _sqlite_backup(DB_PATH, dest)
        except PermissionError:
            folder = os.path.join(os.path.dirname(DB_PATH), "shop_backup")
            os.makedirs(folder, exist_ok=True)
            dest = os.path.join(folder, f"shop_{timestamp}.db")
            _sqlite_backup(DB_PATH, dest)

        backups = sorted(
            [f for f in os.listdir(folder) if f.endswith(".db")], reverse=True
        )
        for old in backups[30:]:
            os.remove(os.path.join(folder, old))

        try:
            _export_excel(folder, timestamp)
        except Exception:
            pass

        xlsx_backups = sorted(
            [f for f in os.listdir(folder) if f.endswith(".xlsx")], reverse=True
        )
        for old in xlsx_backups[30:]:
            os.remove(os.path.join(folder, old))

        save_setting("last_backup", datetime.now().strftime("%Y-%m-%d %H:%M"))
        return {"success": True, "path": dest}
    except Exception as e:
        return {"success": False, "error": str(e)}


def auto_backup():
    import services.google_drive as gdb

    if not _backup_lock.acquire(blocking=False):
        return

    try:
        for attempt in range(3):
            try:
                if gdb.get_credentials():
                    ok, msg = gdb.backup_to_drive(DB_PATH)
                    if ok:
                        save_setting(
                            "last_backup",
                            datetime.now().strftime("%Y-%m-%d %H:%M") + " (Drive)",
                        )
                        return
            except Exception:
                if attempt < 2:
                    time.sleep(2)
                    continue
                break

        for attempt in range(3):
            result = do_backup()
            if result["success"]:
                return
            if result.get("error") == "لم يتم تحديد مسار النسخ الاحتياطي":
                return
            if attempt < 2:
                time.sleep(2)
            else:
                print(f"[Backup] فشل: {result.get('error')}")
    finally:
        _backup_lock.release()


def _run_scheduler():
    import schedule
    schedule.every(15).minutes.do(auto_backup)
    while True:
        schedule.run_pending()
        time.sleep(1)


def start_auto_backup():
    global _scheduler_started
    if not _scheduler_started:
        _scheduler_started = True
        t = threading.Thread(target=_run_scheduler, daemon=True)
        t.start()
