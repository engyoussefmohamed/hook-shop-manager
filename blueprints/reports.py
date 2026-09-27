import io
import datetime

from flask import Blueprint, render_template, request, jsonify

from database import (
    get_daily_stats, get_sales_by_date, get_top_products,
    get_last_30_days, get_payment_summary, search_sales,
    get_all_sales_for_export, get_settings, save_setting,
)
from database.reports import get_all_receipts, get_receipts_stats, get_period_stats
from helpers import admin_required, _positive_int
from config import SALES_PER_PAGE

reports_bp = Blueprint("reports", __name__)


@reports_bp.route("/reports")
@admin_required
def reports():
    date       = request.args.get("date", "")
    sale_id    = request.args.get("sale_id", "").strip()
    date_from  = request.args.get("from", "").strip()
    date_to    = request.args.get("to", "").strip()
    payment    = request.args.get("payment", "").strip()
    product_q  = request.args.get("product", "").strip()
    page       = _positive_int(request.args.get("page", 1))
    stats      = get_daily_stats(date or None)
    has_advanced = any([sale_id, date_from, date_to, payment, product_q])
    if has_advanced:
        sales = search_sales(
            sale_id or None, date_from or None, date_to or None,
            payment or None, product_q or None,
        )
        total_sales = len(sales)
        total_pages = 1
    else:
        sales, total_sales = get_sales_by_date(
            date or None, page=page, per_page=SALES_PER_PAGE
        )
        total_pages = max(1, (total_sales + SALES_PER_PAGE - 1) // SALES_PER_PAGE)
    top        = get_top_products(10)
    chart_data = get_last_30_days()
    payments   = get_payment_summary(date or None)

    settings        = get_settings()
    period_start    = settings.get("period_start_date") or \
                      datetime.date.today().replace(day=1).isoformat()
    period_stats    = get_period_stats(period_start)

    return render_template(
        "reports.html", stats=stats, sales=sales,
        top_products=top, chart_data=chart_data,
        selected_date=date, page=page,
        total_pages=total_pages, total_sales=total_sales,
        payments=payments,
        period_stats=period_stats, period_start=period_start,
        filters={
            "sale_id": sale_id, "from": date_from, "to": date_to,
            "payment": payment, "product": product_q,
            "advanced": has_advanced,
        },
    )


@reports_bp.route("/api/reports/reset-period", methods=["POST"])
@admin_required
def reset_period():
    save_setting("period_start_date", datetime.date.today().isoformat())
    return jsonify({"success": True, "date": datetime.date.today().isoformat()})


@reports_bp.route("/reports/export")
@admin_required
def reports_export():
    from services.excel import style_header_row, auto_width, make_xlsx_response
    from openpyxl import Workbook
    from openpyxl.styles import Font

    date_from = request.args.get("from", "")
    date_to   = request.args.get("to", "")

    sales    = get_all_sales_for_export(date_from or None, date_to or None)
    settings = get_settings()
    shop_name = settings.get("shop_name", "محل الملابس")

    wb = Workbook()
    ws = wb.active
    ws.title = "المبيعات"
    ws.sheet_view.rightToLeft = True

    ws.append([shop_name])
    ws["A1"].font = Font(bold=True, size=14)
    label = "تقرير المبيعات"
    if date_from or date_to:
        label += f" ({date_from or '...'} — {date_to or '...'})"
    ws.append([label])
    ws.append([])

    headers = [
        "#", "التاريخ", "الإجمالي", "الخصم", "المدفوع", "الربح",
        "طريقة الدفع", "ملاحظات",
    ]
    ws.append(headers)
    style_header_row(ws, ws.max_row, len(headers))

    total_rev = total_disc = total_paid = total_prof = 0
    for s in sales:
        ws.append([
            s["id"], s["created_at"],
            round(s["total"], 2), round(s["discount"], 2),
            round(s["final_total"], 2), round(s["profit"], 2),
            s["payment_method"], s.get("notes", ""),
        ])
        total_rev  += s["total"]
        total_disc += s["discount"]
        total_paid += s["final_total"]
        total_prof += s["profit"]

    ws.append([])
    ws.append([
        "الإجمالي", "", round(total_rev, 2), round(total_disc, 2),
        round(total_paid, 2), round(total_prof, 2),
    ])
    last = ws.max_row
    for col in [1, 3, 4, 5, 6]:
        ws.cell(row=last, column=col).font = Font(bold=True)

    auto_width(ws)
    fname = f"sales_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return make_xlsx_response(buf, fname)


@reports_bp.route("/receipts")
@admin_required
def receipts_page():
    page         = _positive_int(request.args.get("page", 1))
    date         = request.args.get("date", "").strip()
    payment      = request.args.get("payment", "").strip()
    sale_id      = request.args.get("sale_id", "").strip()

    sales, total_sales = get_all_receipts(
        page=page, per_page=SALES_PER_PAGE,
        date_from=date or None, date_to=date or None,
        payment_method=payment or None, sale_id=sale_id or None,
        show_voided=False,
    )
    total_pages = max(1, -(-total_sales // SALES_PER_PAGE))
    stats = get_receipts_stats(
        date_from=date or None, date_to=date or None,
        show_voided=False,
    )

    return render_template(
        "receipts.html", sales=sales, stats=stats,
        page=page, total_pages=total_pages, total_sales=total_sales,
        filters={
            "date": date,
            "payment": payment, "sale_id": sale_id,
        },
    )
