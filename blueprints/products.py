import io
import os
import sys

import barcode as barcode_lib
from barcode.writer import ImageWriter
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file

from database import (
    get_all_products, get_product_by_id,
    add_product, update_product, delete_product, get_categories,
)
from helpers import admin_required, generate_qr, _positive_int, QR_DIR
from config import PER_PAGE

products_bp = Blueprint("products", __name__)


@products_bp.route("/products")
@admin_required
def products():
    search   = request.args.get("q", "").strip()
    category = request.args.get("category", "")
    page     = _positive_int(request.args.get("page", 1))
    low_only = request.args.get("low") == "1"

    if low_only:
        from database import get_low_stock_custom
        all_low = get_low_stock_custom()
        if search:
            all_low = [p for p in all_low if search.lower() in p["name"].lower()
                       or search in (p.get("barcode") or "")]
        if category:
            all_low = [p for p in all_low if p["category"] == category]
        total = len(all_low)
        offset = (page - 1) * PER_PAGE
        items = all_low[offset: offset + PER_PAGE]
    else:
        items, total = get_all_products(
            search or None, category or None, page=page, per_page=PER_PAGE
        )
    categories  = get_categories()
    total_pages = max(1, (total + PER_PAGE - 1) // PER_PAGE)

    return render_template(
        "products.html", products=items, categories=categories,
        search=search, selected_category=category,
        page=page, total_pages=total_pages, total=total, low_only=low_only,
    )


@products_bp.route("/products/add", methods=["GET", "POST"])
@admin_required
def products_add():
    categories = get_categories()
    if request.method == "POST":
        name       = request.form.get("name", "").strip()
        category   = request.form.get("category", "ملابس")
        price      = request.form.get("price", 0)
        cost_price = request.form.get("cost_price", 0)
        size       = request.form.get("size", "").strip()
        color      = request.form.get("color", "").strip()
        quantity   = request.form.get("quantity", 0)
        barcode    = request.form.get("barcode", "").strip()
        min_stock  = request.form.get("min_stock", 5)

        if not name or not price:
            flash("الاسم والسعر مطلوبان", "danger")
            return render_template(
                "products_form.html", categories=categories, product=None
            )

        try:
            pid = add_product(
                name, category, price, cost_price, size, color, quantity, barcode, min_stock
            )
            generate_qr(pid)
            flash(f"تم إضافة {name} بنجاح ✓", "success")
            return redirect(url_for("products.products"))
        except Exception as e:
            flash(f"خطأ: {e}", "danger")

    return render_template("products_form.html", categories=categories, product=None)


@products_bp.route("/products/<int:pid>/edit", methods=["GET", "POST"])
@admin_required
def products_edit(pid):
    product    = get_product_by_id(pid)
    categories = get_categories()

    if not product:
        flash("المنتج غير موجود", "danger")
        return redirect(url_for("products.products"))

    if request.method == "POST":
        update_product(
            pid,
            request.form.get("name", "").strip(),
            request.form.get("category", "ملابس"),
            request.form.get("price", 0),
            request.form.get("cost_price", 0),
            request.form.get("size", "").strip(),
            request.form.get("color", "").strip(),
            request.form.get("quantity", 0),
            request.form.get("barcode", "").strip() or product["barcode"],
            request.form.get("min_stock", 5),
        )
        generate_qr(pid)
        flash("تم التعديل بنجاح ✓", "success")
        return redirect(url_for("products.products"))

    return render_template(
        "products_form.html", categories=categories, product=product
    )


@products_bp.route("/products/<int:pid>/delete", methods=["POST"])
@admin_required
def products_delete(pid):
    product = get_product_by_id(pid)
    if product:
        delete_product(pid)
        qr_path = os.path.join(QR_DIR, f"{pid}.png")
        if os.path.exists(qr_path):
            os.remove(qr_path)
        flash(f"تم حذف {product['name']} ✓", "success")
    return redirect(url_for("products.products"))


@products_bp.route("/products/<int:pid>/qr")
@admin_required
def products_qr(pid):
    product = get_product_by_id(pid)
    if not product:
        return "المنتج غير موجود", 404

    qr_path = os.path.join(QR_DIR, f"{pid}.png")
    if not os.path.exists(qr_path):
        generate_qr(pid)

    return send_file(qr_path, mimetype="image/png")


@products_bp.route("/products/<int:pid>/barcode")
@admin_required
def products_barcode(pid):
    product = get_product_by_id(pid)
    if not product:
        return "المنتج غير موجود", 404

    buf = io.BytesIO()
    code_value = product["barcode"] or str(product["id"])
    writer = ImageWriter()
    if getattr(sys, "frozen", False):
        font_path = os.path.join(sys._MEIPASS, "barcode", "fonts", "DejaVuSansMono.ttf")
        if os.path.exists(font_path):
            writer.font_path = font_path
    code = barcode_lib.get("code128", code_value, writer=writer)
    code.write(buf, options={
        "module_height": 12.0,
        "module_width":  0.25,
        "quiet_zone":    4.0,
        "font_size":     8,
        "text_distance": 3.0,
        "background":    "white",
        "foreground":    "black",
        "write_text":    True,
    })
    buf.seek(0)
    return send_file(buf, mimetype="image/png")


@products_bp.route("/products/export")
@admin_required
def products_export():
    from services.excel import style_header_row, auto_width, make_xlsx_response
    from openpyxl import Workbook
    from openpyxl.styles import Font

    products = get_all_products()
    wb = Workbook()
    ws = wb.active
    ws.title = "المنتجات"
    ws.sheet_view.rightToLeft = True
    headers = [
        "الاسم", "الفئة", "سعر البيع", "سعر التكلفة", "الكمية",
        "الباركود", "حد أدنى",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))
    for p in products:
        ws.append([
            p["name"], p["category"], p["price"], p["cost_price"],
            p["quantity"], p["barcode"],
            p.get("min_stock", 5),
        ])
    auto_width(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return make_xlsx_response(buf, "products.xlsx")


@products_bp.route("/products/bulk-delete", methods=["POST"])
@admin_required
def products_bulk_delete():
    ids   = request.form.getlist("ids")
    count = 0
    for pid_str in ids:
        try:
            pid = int(pid_str)
            product = get_product_by_id(pid)
            if product:
                delete_product(pid)
                qr_path = os.path.join(QR_DIR, f"{pid}.png")
                if os.path.exists(qr_path):
                    os.remove(qr_path)
                count += 1
        except (ValueError, TypeError):
            pass
    flash(f"تم حذف {count} منتج ✓", "success")
    return redirect(url_for("products.products"))


@products_bp.route("/products/import", methods=["POST"])
@admin_required
def products_import():
    from database import get_product_by_barcode as _gpb
    from database.settings import add_category
    file = request.files.get("file")
    if not file:
        flash("اختر ملف Excel أولاً", "danger")
        return redirect(url_for("products.products"))
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, data_only=True)
        ws = wb.active
        headers = [str(c.value or "").strip() for c in ws[1]]
        aliases = {
            "name":       ["name", "اسم", "الاسم", "اسم المنتج"],
            "category":   ["category", "فئة", "الفئة"],
            "price":      ["price", "سعر", "سعر البيع"],
            "cost_price": ["cost_price", "تكلفة", "سعر التكلفة"],
            "quantity":   ["quantity", "كمية", "الكمية"],
            "barcode":    ["barcode", "باركود", "الباركود"],
            "min_stock":  ["min_stock", "حد أدنى", "الحد الأدنى"],
        }
        index = {}
        for key, names in aliases.items():
            for idx, title in enumerate(headers):
                if title.lower() in [n.lower() for n in names]:
                    index[key] = idx
                    break
        if "name" not in index or "price" not in index:
            flash("ملف Excel لازم يحتوي الاسم وسعر البيع على الأقل", "danger")
            return redirect(url_for("products.products"))
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            def val(key, default=""):
                pos = index.get(key)
                return row[pos] if pos is not None and pos < len(row) and row[pos] is not None else default
            name = str(val("name", "")).strip()
            if not name:
                continue
            barcode_val = str(val("barcode", "")).strip()
            existing = _gpb(barcode_val) if barcode_val else None
            payload = {
                "name": name,
                "category": str(val("category", "ملابس")).strip() or "ملابس",
                "price": float(val("price", 0) or 0),
                "cost_price": float(val("cost_price", 0) or 0),
                "size": "",
                "color": "",
                "quantity": int(val("quantity", 0) or 0),
                "barcode": barcode_val,
                "min_stock": int(val("min_stock", 5) or 5),
            }
            add_category(payload["category"])
            if existing:
                # Update info only — keep current stock quantity untouched
                update_product(
                    existing["id"],
                    name       = payload["name"],
                    category   = payload["category"],
                    price      = payload["price"],
                    cost_price = payload["cost_price"],
                    size       = existing.get("size", ""),
                    color      = existing.get("color", ""),
                    quantity   = existing["quantity"],
                    barcode    = payload["barcode"],
                    min_stock  = payload["min_stock"],
                )
            else:
                add_product(
                    payload["name"], payload["category"], payload["price"],
                    payload["cost_price"], payload["size"], payload["color"],
                    payload["quantity"], payload["barcode"], payload["min_stock"],
                )
            count += 1
        flash(f"تم استيراد/تحديث {count} منتج", "success")
    except Exception as exc:
        flash(f"تعذر الاستيراد: {exc}", "danger")
    return redirect(url_for("products.products"))
